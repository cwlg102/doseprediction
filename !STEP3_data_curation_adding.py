import os 
import SimpleITK as sitk
import openpyxl 
import numpy as np

xl_path = r"X:\!project\!2025_doseprediction\abdominal_dose_patients_cl_jk_update.xlsx"

wb = openpyxl.load_workbook(xl_path)
ws = wb['Sheet']

savepath_site = r"X:\!project\!2025_doseprediction\data\processed_data\traindata_128_from800downsample_add_pres_chan\imagesTr_additional_info_site"
savepath_treat = r"X:\!project\!2025_doseprediction\data\processed_data\traindata_128_from800downsample_add_pres_chan\imagesTr_additional_info_treatment"
os.makedirs(savepath_site, exist_ok=True)
os.makedirs(savepath_treat, exist_ok=True)
im_basepath = r"X:\!project\!2025_doseprediction\data\processed_data\traindata_128_from800downsample_add_pres_chan\imagesTr_MR"
im_dirs_list = sorted(os.listdir(im_basepath))

for idx, im_dir in enumerate(im_dirs_list):
    if str(ws.cell(idx + 3, 1).value) == "1642725":
        continue
    im_itk = sitk.ReadImage(os.path.join(im_basepath, im_dir))
    im_arr = sitk.GetArrayFromImage(im_itk)
    site_arr = np.zeros_like(im_arr)
    
    if str(ws.cell(idx + 3, 3).value) == "Liver":
        site_arr += 1
    elif str(ws.cell(idx + 3, 3).value) == "Pancreas":
        site_arr += 2
    else:
        site_arr += 3 
    site_arr = np.uint8(site_arr)
    
    site_itk = sitk.GetImageFromArray(site_arr)
    site_itk.CopyInformation(im_itk)
    
    treat_arr = np.zeros_like(im_arr)
    
    if "SBRT" in str(ws.cell(idx + 3, 4).value) or "%" in str(ws.cell(idx + 3, 4).value):
        treat_arr += 1 
    else:
        pass
    treat_arr = np.uint8(treat_arr)
    
    treat_itk = sitk.GetImageFromArray(treat_arr)
    treat_itk.CopyInformation(im_itk)
    
    sitk.WriteImage(site_itk, os.path.join(savepath_site, im_dir.split("_")[0] + "_" + im_dir.split("_")[1] + "_site.nii.gz"))
    sitk.WriteImage(treat_itk, os.path.join(savepath_treat, im_dir.split("_")[0] + "_" + im_dir.split("_")[1] + "_treat.nii.gz"))
